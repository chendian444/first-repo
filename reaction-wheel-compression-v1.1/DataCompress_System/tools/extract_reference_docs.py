from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document


def clean_line(value: str) -> str:
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    lines: list[str] = [f"# {path.name}", ""]
    for para in doc.paragraphs:
        text = clean_line(para.text)
        if text:
            lines.append(text)
    for table_index, table in enumerate(doc.tables, 1):
        lines.append("")
        lines.append(f"## TABLE {table_index}")
        for row in table.rows:
            cells = [clean_line(cell.text) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_binary_doc(path: Path) -> str:
    data = path.read_bytes()
    candidates: list[str] = []
    for encoding in ("utf-16le", "gb18030", "latin1"):
        try:
            text = data.decode(encoding, errors="ignore")
        except LookupError:
            continue
        text = text.replace("\r", "\n")
        for part in re.split(r"[\n\t\x00]+", text):
            part = clean_line(part)
            chinese_count = len(re.findall(r"[\u4e00-\u9fff]", part))
            useful_chars = len(re.findall(r"[\u4e00-\u9fffA-Za-z0-9（）()：:、，,。.；;/%._-]", part))
            if len(part) >= 4 and chinese_count >= 2 and useful_chars / max(1, len(part)) > 0.45:
                candidates.append(part)
    deduped: list[str] = []
    seen: set[str] = set()
    for item in candidates:
        if item not in seen:
            seen.add(item)
            deduped.append(item)
    return "\n".join([f"# {path.name}", "", *deduped])


def main() -> int:
    if len(sys.argv) < 3:
        print("usage: extract_reference_docs.py <out-dir> <doc-path> [<doc-path> ...]")
        return 2
    out_dir = Path(sys.argv[1])
    out_dir.mkdir(parents=True, exist_ok=True)
    for raw in sys.argv[2:]:
        path = Path(raw)
        if path.suffix.lower() == ".docx":
            content = extract_docx(path)
        else:
            content = extract_binary_doc(path)
        target = out_dir / f"{path.stem}.txt"
        target.write_text(content, encoding="utf-8")
        print(target)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
